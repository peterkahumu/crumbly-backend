import csv
import io
from typing import Any, Dict, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


# ──────────────────────────────────────────────── Legacy / Basic Helpers


def export_to_csv(headers: List[str], data: List[Dict[str, Any]]) -> io.BytesIO:
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=headers)
    writer.writeheader()
    for row in data:
        writer.writerow({h: row.get(h, "") for h in headers})

    bytes_io = io.BytesIO(output.getvalue().encode("utf-8"))
    return bytes_io


def export_to_excel(headers: List[str], data: List[Dict[str, Any]]) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)

    for row_data in data:
        row = [str(row_data.get(h, "")) for h in headers]
        ws.append(row)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def export_to_pdf(title: str, headers: List[str], data: List[Dict[str, Any]]) -> io.BytesIO:
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(letter))
    elements = []

    styles = getSampleStyleSheet()
    elements.append(Paragraph(title, styles["Title"]))

    table_data = [headers]
    for row in data:
        table_data.append([str(row.get(h, "")) for h in headers])

    table = Table(table_data)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
                ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
                ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ]
        )
    )

    elements.append(table)
    doc.build(elements)
    output.seek(0)
    return output


# ──────────────────────────────────────────────── Rich Stats / Analytics Exporters


class NumberedCanvas(canvas.Canvas):
    """Custom canvas that adds page numbers ('Page X of Y') and footers to every page."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#6B4038"))
        self.setStrokeColor(colors.HexColor("#E5D9D3"))
        self.setLineWidth(0.5)

        # Footer divider and text
        self.line(36, 30, A4[0] - 36, 30)
        self.drawString(36, 18, "Crumbly Marketplace — Performance Report")
        self.drawRightString(A4[0] - 36, 18, f"Page {self._pageNumber} of {page_count}")
        self.restoreState()


def build_analytics_workbook(report: dict) -> io.BytesIO:
    """Builds a rich multi-sheet Excel workbook from an analytics report dict."""
    wb = openpyxl.Workbook()
    # Remove initial default sheet
    wb.remove(wb.active)

    brand_fill = PatternFill(start_color="C1502E", end_color="C1502E", fill_type="solid")
    alt_fill = PatternFill(start_color="FBF7F0", end_color="FBF7F0", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="E5D9D3"),
        right=Side(style="thin", color="E5D9D3"),
        top=Side(style="thin", color="E5D9D3"),
        bottom=Side(style="thin", color="E5D9D3"),
    )

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10, color="2E1B19")

    for sheet_data in report.get("sheets", []):
        sheet_name = sheet_data["name"][:31]
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True

        cols = sheet_data.get("columns", [])
        num_cols = max(len(cols), 2)
        money_cols = set(sheet_data.get("money_cols", []))

        # Row 1: Title
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
        c1 = ws.cell(row=1, column=1)
        c1.value = f"{report.get('title', 'Crumbly')} — {sheet_data['name']}"
        c1.font = Font(name="Calibri", size=13, bold=True, color="2E1B19")
        ws.row_dimensions[1].height = 24

        # Row 2: Subtitle
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols)
        c2 = ws.cell(row=2, column=1)
        period = report.get("period_label", "")
        drange = report.get("date_range", "")
        gen_at = report.get("generated_at", "")
        c2.value = f"{period} ({drange}) · Generated {gen_at}"
        c2.font = Font(name="Calibri", size=9, italic=True, color="6B4038")
        ws.row_dimensions[2].height = 18

        # Row 3: Spacing
        ws.row_dimensions[3].height = 8

        # Row 4: Column Headers
        ws.row_dimensions[4].height = 24
        for col_idx, col_name in enumerate(cols, start=1):
            cell = ws.cell(row=4, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = brand_fill
            cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
            cell.border = thin_border

        # Row 5+: Data rows
        rows = sheet_data.get("rows", [])
        for r_idx, row_values in enumerate(rows, start=5):
            ws.row_dimensions[r_idx].height = 20
            fill = alt_fill if r_idx % 2 == 0 else white_fill
            for c_idx, val in enumerate(row_values, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = data_font
                cell.fill = fill
                cell.border = thin_border

                col_idx_zero_based = c_idx - 1
                if col_idx_zero_based in money_cols and isinstance(val, (int, float)):
                    cell.number_format = '#,##0 "KES"'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif isinstance(val, (int, float)):
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-adjust column widths
        for col_idx in range(1, len(cols) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = len(str(cols[col_idx - 1]))
            for r in rows:
                if col_idx - 1 < len(r):
                    val_str = str(r[col_idx - 1])
                    if len(val_str) > max_len:
                        max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 46)

        # Freeze header row
        ws.freeze_panes = "A5"

        # Auto-filter
        if cols and rows:
            ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{4 + len(rows)}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_analytics_pdf(report: dict) -> io.BytesIO:
    """Builds a formatted, multi-page PDF from an analytics report dict."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=45,
    )

    page_width = A4[0] - 72  # 523.27 pt

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CoverTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.white,
    )
    subtitle_style = ParagraphStyle(
        "CoverSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#FBE2D6"),
    )
    section_title_style = ParagraphStyle(
        "SectionTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#C1502E"),
        spaceBefore=14,
        spaceAfter=6,
    )
    meta_text_style = ParagraphStyle(
        "MetaText",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#6B4038"),
    )
    cell_style = ParagraphStyle(
        "CellText",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#2E1B19"),
    )
    cell_right_style = ParagraphStyle(
        "CellRight",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#2E1B19"),
        alignment=2,
    )
    cell_header_style = ParagraphStyle(
        "CellHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
    )

    elements = []

    # 1. Header Banner
    header_content = [
        [Paragraph(report.get("title", "Crumbly Marketplace"), title_style)],
        [Paragraph(report.get("subtitle", "Performance Report"), subtitle_style)],
    ]
    header_table = Table(header_content, colWidths=[page_width])
    header_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#C1502E")),
                ("TOPPADDING", (0, 0), (-1, 0), 12),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
                ("LEFTPADDING", (0, 0), (-1, -1), 16),
                ("RIGHTPADDING", (0, 0), (-1, -1), 16),
            ]
        )
    )
    elements.append(header_table)
    elements.append(Spacer(1, 10))

    # 2. Meta Info Box
    filters_summary = " · ".join([f"{k}: {v}" for k, v in report.get("filters", [])])
    meta_lines = [
        f"<b>Period:</b> {report.get('period_label', '')} ({report.get('date_range', '')})    ·    <b>Generated:</b> {report.get('generated_at', '')}"
    ]
    if filters_summary:
        meta_lines.append(f"<b>Applied Filters:</b> {filters_summary}")

    meta_table = Table([[Paragraph("<br/>".join(meta_lines), meta_text_style)]], colWidths=[page_width])
    meta_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FBF7F0")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5D9D3")),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("LEFTPADDING", (0, 0), (-1, -1), 12),
                ("RIGHTPADDING", (0, 0), (-1, -1), 12),
            ]
        )
    )
    elements.append(meta_table)
    elements.append(Spacer(1, 12))

    # 3. KPI Tiles
    kpis = report.get("kpis", [])
    if kpis:
        kpi_per_row = 3
        kpi_rows = []
        for i in range(0, len(kpis), kpi_per_row):
            chunk = kpis[i : i + kpi_per_row]
            row_cells = []
            for item in chunk:
                cell_html = (
                    f"<font size='8' color='#6B4038'>{item['label'].upper()}</font><br/>"
                    f"<font size='13' color='#2E1B19'><b>{item['value']}</b></font>"
                )
                row_cells.append(Paragraph(cell_html, styles["Normal"]))
            while len(row_cells) < kpi_per_row:
                row_cells.append("")
            kpi_rows.append(row_cells)

        col_w = page_width / kpi_per_row
        kpi_table = Table(kpi_rows, colWidths=[col_w] * kpi_per_row)
        kpi_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFFFFF")),
                    ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5D9D3")),
                    ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5D9D3")),
                    ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                    ("LEFTPADDING", (0, 0), (-1, -1), 10),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ]
            )
        )
        elements.append(kpi_table)
        elements.append(Spacer(1, 14))

    # 4. Sheet Tables
    for sheet in report.get("sheets", []):
        cols = sheet.get("columns", [])
        rows = sheet.get("rows", [])
        if not cols or not rows:
            continue

        elements.append(Paragraph(sheet["name"], section_title_style))

        money_cols = set(sheet.get("money_cols", []))
        table_data = []

        header_row = [Paragraph(f"<b>{c}</b>", cell_header_style) for c in cols]
        table_data.append(header_row)

        display_rows = rows[:40]
        for r in display_rows:
            row_cells = []
            for c_idx, val in enumerate(r):
                if c_idx in money_cols and isinstance(val, (int, float)):
                    txt = f"KES {val:,.0f}"
                    row_cells.append(Paragraph(txt, cell_right_style))
                elif isinstance(val, (int, float)):
                    txt = f"{val:,.0f}"
                    row_cells.append(Paragraph(txt, cell_right_style))
                else:
                    txt = str(val) if val is not None else "—"
                    align = (
                        cell_right_style
                        if c_idx > 0 and (isinstance(val, (int, float)) or txt.endswith("%"))
                        else cell_style
                    )
                    row_cells.append(Paragraph(txt, align))
            table_data.append(row_cells)

        num_c = len(cols)
        first_w = max(page_width * 0.35, 120) if num_c > 2 else page_width * 0.5
        rem_w = (page_width - first_w) / max(num_c - 1, 1)
        col_widths = [first_w] + [rem_w] * (num_c - 1)

        t = Table(table_data, colWidths=col_widths, repeatRows=1)
        t_style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#C1502E")),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5D9D3")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5D9D3")),
        ]
        for row_idx in range(1, len(table_data)):
            bg = colors.HexColor("#FBF7F0") if row_idx % 2 == 0 else colors.white
            t_style.append(("BACKGROUND", (0, row_idx), (-1, row_idx), bg))

        t.setStyle(TableStyle(t_style))
        elements.append(t)
        elements.append(Spacer(1, 10))

    doc.build(elements, canvasmaker=NumberedCanvas)
    buf.seek(0)
    return buf


def build_analytics_csv(report: dict) -> io.BytesIO:
    """Builds a formatted CSV text buffer from an analytics report dict."""
    output = io.StringIO()
    writer = csv.writer(output)

    # Title & Metadata
    writer.writerow([report.get("title", "Crumbly Marketplace")])
    writer.writerow([report.get("subtitle", "Performance Report")])
    writer.writerow(
        [
            f"Period: {report.get('period_label', '')} ({report.get('date_range', '')})",
            f"Generated: {report.get('generated_at', '')}",
        ]
    )
    if report.get("filters"):
        writer.writerow(["Filters", ", ".join(f"{k}: {v}" for k, v in report.get("filters", []))])
    writer.writerow([])

    # KPIs
    writer.writerow(["--- KEY PERFORMANCE INDICATORS ---"])
    for kpi in report.get("kpis", []):
        writer.writerow([kpi["label"], kpi["value"]])
    writer.writerow([])

    # Sheets
    for sheet in report.get("sheets", []):
        writer.writerow([f"--- {sheet['name'].upper()} ---"])
        writer.writerow(sheet.get("columns", []))
        for row in sheet.get("rows", []):
            writer.writerow(row)
        writer.writerow([])

    return io.BytesIO(output.getvalue().encode("utf-8"))
